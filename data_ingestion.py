import os
from supabase import create_client, Client
from pypdf import PdfReader
from docx import Document
import openpyxl
from sentence_transformers import SentenceTransformer
from dotenv import load_dotenv

# 1. HARD-LOCK PATHS (Prevents amnesia)
load_dotenv()
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LIBRARY_DIR = os.path.join(BASE_DIR, "library")

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

print(f"--- Data Ingestion ---")
print(f"Reading from: {LIBRARY_DIR}")
print(f"Saving to: Supabase Cloud")

# 2. Initialize Supabase and Model
supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)
model = SentenceTransformer('all-MiniLM-L6-v2')

# We selectively delete existing file chunks but LEAVE manual entries intact
try:
    response = supabase.table("liaison_library").select("id, metadata").execute()
    ids_to_delete = [row["id"] for row in response.data if not row["metadata"].get("source", "").startswith("Contributor:")]
    if ids_to_delete:
        for chunk_id in ids_to_delete:
            supabase.table("liaison_library").delete().eq("id", chunk_id).execute()
except Exception as e:
    print(f"Notice: Table clear error or empty table - {e}")

def process_pdfs():
    if not os.path.exists(LIBRARY_DIR):
        print(f"ERROR: Library folder not found at {LIBRARY_DIR}")
        return

    pdf_files = [f for f in os.listdir(LIBRARY_DIR) if f.endswith('.pdf')]
    
    if not pdf_files:
        print("No PDF files found in the library folder.")
        return

    for filename in pdf_files:
        print(f"Indexing: {filename}...")
        path = os.path.join(LIBRARY_DIR, filename)
        
        try:
            reader = PdfReader(path)
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    text = text.replace('\x00', '') # Remove null bytes unsupported by PostgreSQL
                if text and len(text.strip()) > 50: # Ignore empty/short pages
                    chunk_id = f"{filename}_pg_{i}"
                    embedding = model.encode([text])[0].tolist()
                    
                    supabase.table("liaison_library").upsert({
                        "id": chunk_id,
                        "embedding": embedding,
                        "content": text,
                        "metadata": {"source": filename, "page": i + 1}
                    }).execute()
            print(f"Done: {filename}")
        except Exception as e:
            print(f"Error processing {filename}: {e}")

def process_word_docs():
    if not os.path.exists(LIBRARY_DIR):
        print(f"ERROR: Library folder not found at {LIBRARY_DIR}")
        return

    word_files = [f for f in os.listdir(LIBRARY_DIR) if f.endswith('.docx')]
    
    if not word_files:
        print("No Word documents found in the library folder.")
        return

    for filename in word_files:
        print(f"Indexing: {filename}...")
        path = os.path.join(LIBRARY_DIR, filename)
        
        try:
            doc = Document(path)
            full_text = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    full_text.append(paragraph.text)
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            full_text.append(cell.text)
            
            text = '\n'.join(full_text)
            if text:
                text = text.replace('\x00', '')
            if text and len(text.strip()) > 50:
                chunk_id = f"{filename}_doc"
                embedding = model.encode([text])[0].tolist()
                
                supabase.table("liaison_library").upsert({
                    "id": chunk_id,
                    "embedding": embedding,
                    "content": text,
                    "metadata": {"source": filename, "type": "word"}
                }).execute()
            print(f"Done: {filename}")
        except Exception as e:
            print(f"Error processing {filename}: {e}")

def process_excel_files():
    if not os.path.exists(LIBRARY_DIR):
        print(f"ERROR: Library folder not found at {LIBRARY_DIR}")
        return

    excel_files = [f for f in os.listdir(LIBRARY_DIR) if f.endswith(('.xlsx', '.xls'))]
    
    if not excel_files:
        print("No Excel files found in the library folder.")
        return

    for filename in excel_files:
        print(f"Indexing: {filename}...")
        path = os.path.join(LIBRARY_DIR, filename)
        
        try:
            workbook = openpyxl.load_workbook(path, data_only=True)
            full_text = []
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                full_text.append(f"Sheet: {sheet_name}")
                
                for row in sheet.iter_rows(values_only=True):
                    # Convert all values to strings and filter out None values
                    row_text = [str(cell) for cell in row if cell is not None]
                    if row_text:
                        full_text.append(' | '.join(row_text))
            
            text = '\n'.join(full_text)
            if text:
                text = text.replace('\x00', '')
            if text and len(text.strip()) > 50:
                chunk_id = f"{filename}_xls"
                embedding = model.encode([text])[0].tolist()
                
                supabase.table("liaison_library").upsert({
                    "id": chunk_id,
                    "embedding": embedding,
                    "content": text,
                    "metadata": {"source": filename, "type": "excel"}
                }).execute()
            print(f"Done: {filename}")
        except Exception as e:
            print(f"Error processing {filename}: {e}")

if __name__ == "__main__":
    process_pdfs()
    process_word_docs()
    process_excel_files()
    print(f"--- Ingestion Complete ---")
    response = supabase.table("liaison_library").select("count", count="exact").execute()
    print(f"Total items in AI memory: {response.count}")