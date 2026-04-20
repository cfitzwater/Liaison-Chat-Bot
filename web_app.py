import os
import uuid
from flask import Flask, render_template, request, redirect, url_for, jsonify, send_from_directory, session
from werkzeug.utils import secure_filename
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import generate_password_hash, check_password_hash
from pypdf import PdfReader
from docx import Document
import openpyxl
from urllib.parse import quote_plus
from models import db, User, ChatHistory, LibraryEntry, LibraryFile
from flasgger import Swagger
from dotenv import load_dotenv
from flask_wtf.csrf import CSRFProtect
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address

# Load environment variables from .env file
load_dotenv()

# Link to AI Engine
try:
    from app import search_documents_web, collection, model, CHROMA_PATH
    print(f"--- Web App Linked to AI Memory at: {CHROMA_PATH} ---")
except ImportError:
    print("CRITICAL: app.py not found.")

app = Flask(__name__)
swagger = Swagger(app, template={
    "info": {
        "title": "Liaison Chat Bot API",
        "description": "API documentation for the Liaison Chat Bot and Knowledge Library. **Note:** All API endpoints require active session authentication. You must authenticate via the `/login` web interface to obtain a valid session cookie before making API requests.",
        "version": "1.0.0"
    }
})
app.secret_key = os.getenv('SECRET_KEY', 'liaison_secret_key_123')
app.config['SQLALCHEMY_DATABASE_URI'] = os.getenv('DATABASE_URL', 'sqlite:///project.db')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

# Production Security Settings
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_prefix=1)
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
if os.getenv('FLASK_DEBUG', 'False').lower() not in ('true', '1', 't'):
    app.config['SESSION_COOKIE_SECURE'] = True

ALLOWED_EXTENSIONS = {'pdf', 'docx', 'xlsx', 'xls'}

csrf = CSRFProtect(app)

limiter = Limiter(
    get_remote_address,
    app=app,
    default_limits=["200 per day", "50 per hour"],
    storage_uri="memory://"
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
app.config['UPLOAD_FOLDER'] = os.path.join(BASE_DIR, 'library')

db.init_app(app)

with app.app_context():
    db.create_all()
    
    # Auto-sync files in the library folder to the database
    if os.path.exists(app.config['UPLOAD_FOLDER']):
        for filename in os.listdir(app.config['UPLOAD_FOLDER']):
            if filename.startswith('.'): continue
            if os.path.isfile(os.path.join(app.config['UPLOAD_FOLDER'], filename)):
                if not LibraryFile.query.filter_by(filename=filename).first():
                    db.session.add(LibraryFile(filename=filename))
        db.session.commit()

# --- AUTH & SIGNUP ---
@app.route('/login', methods=['GET', 'POST'])
@limiter.limit("5 per minute")
def login():
    error = None
    if request.method == 'POST':
        email = request.form.get('email').lower()
        password = request.form.get('password')
        user = User.query.filter_by(email=email).first()
        if user and user.is_active:
            if user.password == password or check_password_hash(user.password, password):
                # Migrate plain password to hash if it was plain
                if user.password == password:
                    user.password = generate_password_hash(password)
                    db.session.commit()
                session.update({'user_id': user.id, 'user_email': user.email, 'first_name': user.first_name, 'last_name': user.last_name, 'is_admin': user.is_admin})
                if user.force_password_change:
                    return redirect(url_for('change_password'))
                return redirect(url_for('index'))
            else:
                error = "Invalid email or password."
        else:
            error = "Invalid email or password."
    return render_template('login.html', error=error)

@app.route('/change_password', methods=['GET', 'POST'])
def change_password():
    if not session.get('user_id'):
        return redirect(url_for('login'))
    user = User.query.get(session['user_id'])
    if not user or not user.force_password_change:
        return redirect(url_for('index'))
    if request.method == 'POST':
        new_password = request.form.get('new_password')
        confirm_password = request.form.get('confirm_password')
        if new_password != confirm_password:
            return render_template('change_password.html', error="Passwords do not match.")
        if len(new_password) < 6:
            return render_template('change_password.html', error="Password must be at least 6 characters.")
        user.password = generate_password_hash(new_password)
        user.force_password_change = False
        db.session.commit()
        return redirect(url_for('index'))
    return render_template('change_password.html')

@app.route('/signup', methods=['GET', 'POST'])
@limiter.limit("5 per minute")
def signup():
    if request.method == 'POST':
        email = request.form.get('email').lower()
        existing_user = User.query.filter_by(email=email).first()
        if existing_user:
            if not existing_user.is_active:
                # Reactivate the user
                existing_user.first_name = request.form.get('first_name')
                existing_user.last_name = request.form.get('last_name')
                existing_user.password = generate_password_hash(request.form.get('password'))
                existing_user.is_active = True
                db.session.commit()
                return redirect(url_for('login'))
            else:
                # User already active
                return "User already exists and is active", 400
        else:
            # Create new user
            is_admin = (email == 'fitz3663@gmail.com' or User.query.first() is None)
            new_user = User(
                first_name=request.form.get('first_name'), last_name=request.form.get('last_name'),
                email=email, password=generate_password_hash(request.form.get('password')), is_admin=is_admin
            )
            db.session.add(new_user)
            db.session.commit()
            return redirect(url_for('login'))
    return render_template('signup.html')

@app.route('/logout')
def logout():
    session.clear(); return redirect(url_for('login'))

# --- CHAT & HISTORY ---
@app.route('/')
def index():
    if 'user_id' not in session: return redirect(url_for('login'))
    chats = ChatHistory.query.filter_by(user_id=session['user_id']).order_by(ChatHistory.timestamp.desc()).all()
    seen, history = set(), []
    for c in chats:
        if c.chat_session_id not in seen:
            history.append(c); seen.add(c.chat_session_id)
    files = LibraryFile.query.order_by(LibraryFile.filename).all()
    return render_template('index.html', first_name=session.get('first_name'), last_name=session.get('last_name'), history=history, files=files)

@app.route('/chat', methods=['POST'])
def chat():
    """
    Send a message to the chat bot
    ---
    tags:
      - Chat
    parameters:
      - name: body
        in: body
        required: true
        schema:
          type: object
          properties:
            message:
              type: string
              example: "What is the employee health requirement?"
    responses:
      200:
        description: The bot's response
        schema:
          type: object
          properties:
            answer:
              type: string
              example: "Employees must complete a health screening..."
      400:
        description: Bad Request (missing or invalid message payload)
        schema:
          type: object
          properties:
            answer:
              type: string
              example: "Error: Please provide a valid message."
      401:
        description: Unauthorized (missing or expired session cookie)
    """
    if 'user_id' not in session:
        return jsonify({"answer": "Your session has expired. Please log in again."}), 401
        
    if 'current_chat_id' not in session: session['current_chat_id'] = str(uuid.uuid4())
    
    data = request.get_json(silent=True)
    user_msg = data.get('message') if data else None
    if not user_msg:
        return jsonify({"answer": "Error: Please provide a valid message."}), 400
        
    response = search_documents_web(user_msg)
    new_chat = ChatHistory(user_id=session['user_id'], chat_session_id=session['current_chat_id'], user_message=user_msg, bot_response=response)
    db.session.add(new_chat); db.session.commit()
    return jsonify({"answer": response})

@app.route('/get_session/<session_id>')
def get_session(session_id):
    chats = ChatHistory.query.filter_by(chat_session_id=session_id).order_by(ChatHistory.timestamp.asc()).all()
    session['current_chat_id'] = session_id 
    return jsonify([{"user": c.user_message, "bot": c.bot_response} for c in chats])

@app.route('/new_chat')
def new_chat():
    session.pop('current_chat_id', None)
    return redirect(url_for('index'))

@app.route('/rename_chat', methods=['POST'])
def rename_chat():
    data = request.json
    chat = ChatHistory.query.filter_by(chat_session_id=data.get('session_id')).first()
    if chat:
        chat.user_message = data.get('new_name')
        db.session.commit(); return jsonify({"status": "success"})
    return jsonify({"status": "error"}), 404

@app.route('/delete_chat', methods=['POST'])
def delete_chat():
    ChatHistory.query.filter_by(chat_session_id=request.json.get('session_id')).delete()
    db.session.commit(); return jsonify({"status": "success"})

# --- KNOWLEDGE SYNC ---
@app.route('/add_knowledge')
def add_knowledge():
    if 'user_id' not in session: return redirect(url_for('login'))
    return render_template('add_item.html', full_name=f"{session.get('first_name')} {session.get('last_name')}")

@app.route('/save_item', methods=['POST'])
def save_item():
    entry_id = str(uuid.uuid4())
    notes = request.form.get('notes')
    contributor = request.form.get('user_name')
    label_tag = request.form.get('label_tag')
    db.session.add(LibraryEntry(id=entry_id, notes=notes, label=label_tag, user_id=session['user_id']))
    db.session.commit()
    content = f"LABEL: {label_tag}\nCONTRIBUTOR: {contributor}\nNOTES: {notes}"
    collection.add(ids=[entry_id], embeddings=model.encode([content]).tolist(), documents=[content], metadatas=[{"source": f"Contributor: {contributor} - {label_tag}"}])
    return '<script>window.close();</script>'

@app.route('/my_knowledge')
def my_knowledge():
    if 'user_id' not in session: return redirect(url_for('login'))
    entries = LibraryEntry.query.filter_by(user_id=session['user_id']).all()
    return render_template('my_knowledge.html', entries=entries)

@app.route('/edit_my_knowledge/<entry_id>', methods=['POST'])
def edit_my_knowledge(entry_id):
    entry = db.session.get(LibraryEntry, entry_id)
    if entry and entry.user_id == session.get('user_id'):
        new_label = request.json.get('label')
        new_notes = request.json.get('notes')
        entry.label = new_label
        entry.notes = new_notes
        db.session.commit()
        contributor = f"{session.get('first_name')} {session.get('last_name')}"
        content = f"LABEL: {new_label}\nCONTRIBUTOR: {contributor}\nNOTES: {new_notes}"
        collection.update(ids=[entry_id], embeddings=model.encode([content]).tolist(), documents=[content])
        return jsonify({"status": "success"})
    return jsonify({"status": "denied"}), 403
@app.route('/delete_my_knowledge/<entry_id>', methods=['POST'])
def delete_my_knowledge(entry_id):
    entry = db.session.get(LibraryEntry, entry_id)
    if entry and entry.user_id == session.get('user_id'):
        db.session.delete(entry)
        db.session.commit()
        collection.delete(ids=[entry_id])
        return jsonify({"status": "success"})
    return jsonify({"status": "error"}), 403

# --- API ENDPOINTS ---
@app.route('/api/v1/knowledge')
def api_get_knowledge():
    """
    Get all knowledge entries for the current user
    ---
    tags:
      - Knowledge
    responses:
      200:
        description: A list of knowledge entries
        schema:
          type: array
          items:
            type: object
            properties:
              id:
                type: string
                example: "a03e2ede-68b9-45a7-b2e1-64da1d314e20"
              label:
                type: string
                example: "Project Guidelines"
              notes:
                type: string
                example: "Always use Python 3.10+"
              timestamp:
                type: string
                example: "2026-04-17T12:00:00"
      401:
        description: Unauthorized
    """
    if 'user_id' not in session:
        return jsonify({"error": "Unauthorized"}), 401
    entries = LibraryEntry.query.filter_by(user_id=session['user_id']).all()
    result = []
    for entry in entries:
        result.append({
            "id": entry.id,
            "label": entry.label,
            "notes": entry.notes,
            "timestamp": entry.timestamp.isoformat()
        })
    return jsonify(result)

@app.route('/api/v1/knowledge/<entry_id>')
def api_get_knowledge_item(entry_id):
    """
    Get a specific knowledge entry by ID
    ---
    tags:
      - Knowledge
    parameters:
      - name: entry_id
        in: path
        type: string
        required: true
        description: The ID of the knowledge entry
    responses:
      200:
        description: A single knowledge entry
        schema:
          type: object
          properties:
            id:
              type: string
            label:
              type: string
            notes:
              type: string
            timestamp:
              type: string
      401:
        description: Unauthorized
      404:
        description: Knowledge entry not found
    """
    if 'user_id' not in session:
        return jsonify({"error": "Unauthorized"}), 401
    entry = db.session.get(LibraryEntry, entry_id)
    if not entry or entry.user_id != session['user_id']:
        return jsonify({"error": "Not found"}), 404
    return jsonify({
        "id": entry.id,
        "label": entry.label,
        "notes": entry.notes,
        "timestamp": entry.timestamp.isoformat()
    })

# --- ADMIN PANEL ---
@app.route('/admin')
def admin_dashboard():
    if not session.get('is_admin'): return "Denied", 403
    return render_template('admin.html', users=User.query.filter_by(is_active=True).all(), files=LibraryFile.query.order_by(LibraryFile.filename).all(), entries=LibraryEntry.query.all())

@app.route('/admin/delete_user/<int:user_id>', methods=['POST'])
def admin_delete_user(user_id):
    user = db.session.get(User, user_id)
    if user and user.email != session.get('user_email'):
        user.is_active = False
        db.session.commit()
        return jsonify({"status": "success"})
    return jsonify({"status": "error"}), 400

@app.route('/toggle_admin/<int:user_id>', methods=['POST'])
def toggle_admin(user_id):
    user = db.session.get(User, user_id)
    if user and user.email != session.get('user_email'):
        user.is_admin = not user.is_admin
        db.session.commit(); return jsonify({"status": "success"})
    return jsonify({"status": "error"}), 400

@app.route('/admin/reset_password/<int:user_id>', methods=['POST'])
def admin_reset_password(user_id):
    user = db.session.get(User, user_id)
    if user:
        user.password = generate_password_hash('Liaison1')
        user.force_password_change = True
        db.session.commit()
        return jsonify({"status": "success"})
    return jsonify({"status": "error"}), 400

@app.route('/admin/upload', methods=['POST'])
def admin_upload():
    if not session.get('is_admin'):
        return "Access denied", 403
    
    file = request.files.get('file')
    if file:
        ext = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
        if ext not in ALLOWED_EXTENSIONS:
            return "Invalid file type. Only PDF, DOCX, and Excel files are allowed.", 400
            
        fname = secure_filename(file.filename)
        path = os.path.join(app.config['UPLOAD_FOLDER'], fname)
        
        # Check if file already exists
        existing = LibraryFile.query.filter_by(filename=fname).first()
        if existing:
            # Remove old file from disk
            old_path = os.path.join(app.config['UPLOAD_FOLDER'], fname)
            if os.path.exists(old_path):
                os.remove(old_path)
            # Remove from Chroma
            try:
                collection.delete(where={"source": fname})
            except Exception as e:
                print(f"Error deleting from Chroma: {e}")
            # Remove from db
            db.session.delete(existing)
            db.session.commit()
        
        # Save new file
        file.save(path)
        db.session.add(LibraryFile(filename=fname))
        db.session.commit()
        
        # Ingest the document into Chroma based on file type
        try:
            if fname.lower().endswith('.pdf'):
                # Process PDF
                reader = PdfReader(path)
                for i, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text and len(text.strip()) > 50:  # Ignore empty/short pages
                        chunk_id = f"{fname}_pg_{i}"
                        embedding = model.encode([text]).tolist()
                        collection.add(
                            ids=[chunk_id],
                            embeddings=embedding,
                            documents=[text],
                            metadatas=[{"source": fname, "page": i + 1}]
                        )
            elif fname.lower().endswith('.docx'):
                # Process Word document
                doc = Document(path)
                full_text = []
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip():
                        full_text.append(paragraph.text)
                
                # Also extract text from tables
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                full_text.append(cell.text)
                
                text = '\n'.join(full_text)
                if text and len(text.strip()) > 50:
                    chunk_id = f"{fname}_doc"
                    embedding = model.encode([text]).tolist()
                    collection.add(
                        ids=[chunk_id],
                        embeddings=embedding,
                        documents=[text],
                        metadatas=[{"source": fname, "type": "word"}]
                    )
            elif fname.lower().endswith(('.xlsx', '.xls')):
                # Process Excel file
                workbook = openpyxl.load_workbook(path, data_only=True)
                full_text = []
                
                for sheet_name in workbook.sheetnames:
                    sheet = workbook[sheet_name]
                    full_text.append(f"Sheet: {sheet_name}")
                    
                    for row in sheet.iter_rows(values_only=True):
                        # Convert all values to strings and filter out None values
                        row_text = [str(cell) for cell in row if cell is not None]
                        if row_text:
                            full_text.append(' | '.join(row_text))
                
                text = '\n'.join(full_text)
                if text and len(text.strip()) > 50:
                    chunk_id = f"{fname}_xls"
                    embedding = model.encode([text]).tolist()
                    collection.add(
                        ids=[chunk_id],
                        embeddings=embedding,
                        documents=[text],
                        metadatas=[{"source": fname, "type": "excel"}]
                    )
        except Exception as e:
            print(f"Error ingesting {fname}: {e}")
    
    return redirect(url_for('admin_dashboard'))

@app.route('/admin/delete_file/<int:file_id>', methods=['POST'])
def admin_delete_file(file_id):
    if not session.get('is_admin'):
        return "Access denied", 403
    
    f_rec = db.session.get(LibraryFile, file_id)
    if f_rec:
        path = os.path.join(app.config['UPLOAD_FOLDER'], f_rec.filename)
        if os.path.exists(path): os.remove(path)
        db.session.delete(f_rec)
        db.session.commit()
        # Remove from Chroma
        try:
            collection.delete(where={"source": f_rec.filename})
        except Exception as e:
            print(f"Error deleting from Chroma: {e}")
        return jsonify({"status": "success"})
    return jsonify({"status": "error"})

@app.route('/admin/edit_entry/<entry_id>', methods=['POST'])
def admin_edit_entry(entry_id):
    if not session.get('is_admin'): return jsonify({"status": "denied"}), 403
    entry = db.session.get(LibraryEntry, entry_id)
    if entry:
        new_notes = request.json.get('notes')
        new_label = request.json.get('label')
        entry.notes = new_notes
        entry.label = new_label
        db.session.commit()
        contributor = f"{entry.contributor.first_name} {entry.contributor.last_name}"
        content = f"LABEL: {new_label}\nCONTRIBUTOR: {contributor}\nNOTES: {new_notes}"
        collection.update(ids=[entry_id], embeddings=model.encode([content]).tolist(), documents=[content], metadatas=[{"source": f"Contributor: {contributor} - {new_label}"}])
        return jsonify({"status": "success"})
    return jsonify({"status": "error"}), 404

@app.route('/admin/delete_entry/<entry_id>', methods=['POST'])
def admin_delete_entry(entry_id):
    if not session.get('is_admin'): return jsonify({"status": "denied"}), 403
    entry = db.session.get(LibraryEntry, entry_id)
    if entry:
        db.session.delete(entry)
        db.session.commit()
        collection.delete(ids=[entry_id])
        return jsonify({"status": "success"})
    return jsonify({"status": "error"}), 404

def _extract_docx_text(path):
    try:
        doc = Document(path)
        lines = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                lines.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(cell.text for cell in row.cells if cell.text.strip())
                if row_text:
                    lines.append(row_text)

        return '\n'.join(lines)
    except Exception as e:
        print(f"Error extracting DOCX text: {e}")
        return None


def _extract_excel_text(path):
    try:
        workbook = openpyxl.load_workbook(path, data_only=True)
        lines = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            lines.append(f"Sheet: {sheet_name}")
            for row in sheet.iter_rows(values_only=True):
                row_text = ' | '.join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    lines.append(row_text)
        return '\n'.join(lines)
    except Exception as e:
        print(f"Error extracting Excel text: {e}")
        return None


@app.route('/preview/<path:filename>')
def preview_file(filename):
    if 'user_id' not in session:
        return redirect(url_for('login'))

    abs_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    if not os.path.exists(abs_path):
        return "File not found", 404

    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''

    if ext == 'pdf':
        # Embed PDF in browser frame when possible
        return render_template('preview_file.html', filename=filename, content_type='pdf', text=None)

    if ext == 'docx':
        text = _extract_docx_text(abs_path)
        if text is None:
            text = "Unable to extract text from this document. You can download it instead."
        return render_template('preview_file.html', filename=filename, content_type='text', text=text)

    if ext in ['xlsx', 'xls']:
        text = _extract_excel_text(abs_path)
        if text is None:
            text = "Unable to extract text from this spreadsheet. You can download it instead."
        return render_template('preview_file.html', filename=filename, content_type='text', text=text)

    # Default fallback to download/view through browser
    return redirect(url_for('get_file', filename=filename))


@app.route('/library/<path:filename>')
def get_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename, as_attachment=False)

if __name__ == '__main__':
    debug_mode = os.getenv('FLASK_DEBUG', 'False').lower() in ('true', '1', 't')
    app.run(debug=debug_mode, port=5000)